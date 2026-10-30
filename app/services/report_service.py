import io
import json
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
from sqlalchemy.orm import Session
from sqlalchemy import or_, desc, asc

from app.models.user import User
from app.models.recipe import Recipe
from app.models.appointment import Appointment
from app.models.payment import Payment
from app.models.clinical import ClinicalRecord
from app.models.activity_log import ActivityLog
from app.models.patient_link import PatientNutritionistLink
from app.models.ai_nutrition import NutritionalPlan
from app.models.tenant import Tenant
from app.models.company import Branch
from app.schemas.report import (
    ReportColumnMeta,
    ReportEntityMeta,
    ReportQueryRequest,
    ReportQueryResponse,
    VoiceReportCommandResponse,
    VoiceReportSummaryRequest,
    VoiceReportSummaryResponse,
)
from app.services.gemini_service import GeminiService


class ReportService:

    ENTITIES_CONFIG: Dict[str, Dict[str, Any]] = {
        "patients": {
            "label": "Clientes / Pacientes",
            "description": "Directorio y estado de vinculación de clientes",
            "model": User,
            "columns": [
                {"key": "full_name", "label": "Nombre Completo", "type": "string"},
                {"key": "email", "label": "Correo Electrónico", "type": "string"},
                {"key": "phone", "label": "Teléfono", "type": "string"},
                {"key": "is_active", "label": "Activo", "type": "boolean"},
                {"key": "linked_status", "label": "Estado Vinculación", "type": "string"},
                {"key": "nutritionist_name", "label": "Especialista Asignado", "type": "string"},
                {"key": "created_at", "label": "Fecha Registro", "type": "date"},
            ],
            "default_columns": ["full_name", "email", "phone", "is_active", "linked_status", "created_at"],
        },
        "recipes": {
            "label": "Catálogo de Recetas",
            "description": "Recetas nutricionales con desglose de macronutrientes",
            "model": Recipe,
            "columns": [
                {"key": "title", "label": "Nombre de Receta", "type": "string"},
                {"key": "category", "label": "Categoría", "type": "string"},
                {"key": "calories", "label": "Calorías (kcal)", "type": "number"},
                {"key": "protein", "label": "Proteína (g)", "type": "number"},
                {"key": "carbohydrates", "label": "Carbohidratos (g)", "type": "number"},
                {"key": "fats", "label": "Grasas (g)", "type": "number"},
                {"key": "fiber", "label": "Fibra (g)", "type": "number"},
                {"key": "difficulty", "label": "Dificultad", "type": "string"},
                {"key": "servings", "label": "Porciones", "type": "number"},
                {"key": "created_at", "label": "Fecha Creación", "type": "date"},
            ],
            "default_columns": ["title", "category", "calories", "protein", "carbohydrates", "fats", "difficulty"],
        },
        "appointments": {
            "label": "Gestión de Citas",
            "description": "Historial y agenda de consultas nutricionales",
            "model": Appointment,
            "columns": [
                {"key": "patient_name", "label": "Paciente", "type": "string"},
                {"key": "nutritionist_name", "label": "Especialista", "type": "string"},
                {"key": "scheduled_at", "label": "Fecha y Hora Cita", "type": "date"},
                {"key": "status", "label": "Estado", "type": "string"},
                {"key": "reason", "label": "Motivo", "type": "string"},
                {"key": "cancellation_reason", "label": "Motivo Cancelación", "type": "string"},
                {"key": "created_at", "label": "Fecha Solicitud", "type": "date"},
            ],
            "default_columns": ["patient_name", "nutritionist_name", "scheduled_at", "status", "reason"],
        },
        "payments": {
            "label": "Caja y Cobros (Pagos)",
            "description": "Transacciones de suscripciones y servicios PayPal",
            "model": Payment,
            "columns": [
                {"key": "customer_name", "label": "Cliente", "type": "string"},
                {"key": "concept", "label": "Concepto", "type": "string"},
                {"key": "amount", "label": "Monto", "type": "number"},
                {"key": "currency", "label": "Moneda", "type": "string"},
                {"key": "status", "label": "Estado", "type": "string"},
                {"key": "payment_method", "label": "Método", "type": "string"},
                {"key": "paypal_order_id", "label": "Ref. PayPal", "type": "string"},
                {"key": "created_at", "label": "Fecha Transacción", "type": "date"},
            ],
            "default_columns": ["customer_name", "concept", "amount", "currency", "status", "payment_method", "created_at"],
        },
        "clinical_records": {
            "label": "Fichas Clínicas y Diagnósticos",
            "description": "Evolución clínica, metas y diagnósticos nutricionales",
            "model": ClinicalRecord,
            "columns": [
                {"key": "patient_name", "label": "Paciente", "type": "string"},
                {"key": "nutritionist_name", "label": "Especialista", "type": "string"},
                {"key": "diagnosis", "label": "Diagnóstico Clínico", "type": "string"},
                {"key": "clinical_goals", "label": "Metas Terapéuticas", "type": "string"},
                {"key": "evolution_notes", "label": "Notas de Evolución", "type": "string"},
                {"key": "created_at", "label": "Fecha Asiento", "type": "date"},
                {"key": "updated_at", "label": "Última Actualización", "type": "date"},
            ],
            "default_columns": ["patient_name", "nutritionist_name", "diagnosis", "clinical_goals", "created_at"],
        },
        "activity_logs": {
            "label": "Bitácora de Auditoría",
            "description": "Registro de eventos de seguridad y transacciones del sistema",
            "model": ActivityLog,
            "default_columns": ["user_name", "action", "category", "description", "created_at"],
            "columns": [
                {"key": "user_name", "label": "Usuario", "type": "string"},
                {"key": "user_email", "label": "Correo", "type": "string"},
                {"key": "action", "label": "Acción", "type": "string"},
                {"key": "category", "label": "Categoría", "type": "string"},
                {"key": "description", "label": "Detalles", "type": "string"},
                {"key": "created_at", "label": "Fecha / Hora", "type": "date"},
            ],
        },
    }

    ENTITIES_CONFIG['nutritional_plans'] = {
        'label':'Planes nutricionales', 'description':'Planes generados, aprobados y archivados', 'model':NutritionalPlan,
        'columns':[{'key':key,'label':label,'type':kind} for key,label,kind in [('title','Plan','string'),('patient_name','Paciente','string'),('nutritionist_name','Nutricionista','string'),('status','Estado','string'),('daily_calories','Calorías diarias','number'),('created_at','Fecha de creación','date'),('approved_at','Aprobado el','date')]],
        'default_columns':['title','patient_name','nutritionist_name','status','daily_calories','created_at'],
    }
    for _config in ENTITIES_CONFIG.values():
        _config['columns'] += [{'key':'company_name','label':'Empresa','type':'string'},{'key':'branch_name','label':'Sucursal','type':'string'}]

    @classmethod
    def get_available_entities(cls) -> List[ReportEntityMeta]:
        result = []
        for key, conf in cls.ENTITIES_CONFIG.items():
            cols = [ReportColumnMeta(**c) for c in conf["columns"]]
            result.append(
                ReportEntityMeta(
                    entity=key,
                    label=conf["label"],
                    description=conf["description"],
                    available_columns=cols,
                )
            )
        return result

    @classmethod
    def query_report_data(
        cls,
        db: Session,
        req: ReportQueryRequest,
        current_user: User,
    ) -> ReportQueryResponse:
        from app.services.report_analytics_service import ReportAnalyticsService
        from app.services.audit_policy import is_mutation
        from sqlalchemy.orm import joinedload
        entity_key=req.entity.lower().strip()
        if entity_key not in cls.ENTITIES_CONFIG: raise ValueError('Entidad de reporte no disponible.')
        req=req.model_copy(update={'entity':entity_key})
        if req.group_by or req.aggregate or req.metric!='count':return ReportAnalyticsService.query(db,req,current_user)
        conf=cls.ENTITIES_CONFIG[entity_key];all_cols={c['key']:c for c in conf['columns']}
        selected=req.columns or conf['default_columns']
        if any(key not in all_cols for key in selected):raise ValueError('El reporte contiene columnas que no están disponibles.')
        q,model,branch_col,dt,companies,branches=ReportAnalyticsService.source(db,current_user,req)
        if model is Recipe and not req.status:q=q.filter(Recipe.is_active.is_(True))
        key=req.sort_by or ('scheduled_at' if model is Appointment else 'created_at')
        if key not in {c.name for c in model.__table__.columns}:raise ValueError('La columna de orden no está disponible para este reporte.')
        ordering=getattr(model,key);q=q.order_by(ordering.asc() if req.sort_order=='asc' else ordering.desc())
        total=q.order_by(None).count();items=q.limit(req.limit or 500).all()
        names={c.id:c.name for c in companies};branch_names={b.id:b.name for b in branches}
        ids={getattr(item,k,None) for item in items for k in ('patient_id','nutritionist_id','created_by','user_id')}
        ids|={item.id for item in items if model is User}
        people={u.id:u for u in db.query(User).filter(User.id.in_(ids)).all()} if ids else {}
        links={}
        if model is User:
            for link in db.query(PatientNutritionistLink).filter(PatientNutritionistLink.patient_id.in_([i.id for i in items]),PatientNutritionistLink.tenant_id.in_(names)).order_by(PatientNutritionistLink.created_at.desc()).all():
                if link.patient_id not in links:links[link.patient_id]=link
            nutrition_ids=[l.nutritionist_id for l in links.values() if l.nutritionist_id]
            people.update({u.id:u for u in db.query(User).filter(User.id.in_(nutrition_ids)).all()})
        def person_name(identity):return people[identity].full_name if identity in people else 'Sin asignar'
        def formatted(value):return value.strftime('%Y-%m-%d %H:%M') if isinstance(value,datetime) else value
        rows=[]
        for item in items:
            if model is ActivityLog and not is_mutation(item.action):continue
            row={c.name:formatted(getattr(item,c.name)) for c in model.__table__.columns if c.name not in {'hashed_password','developer_key_hash','ip_address'}}
            if model is User:
                link=links.get(item.id);row.update(linked_status=link.status if link else 'UNLINKED',nutritionist_name=person_name(link.nutritionist_id) if link else 'Sin asignar',is_active='Sí' if item.is_active else 'No')
            else:
                row['patient_name']=person_name(getattr(item,'patient_id',None));row['nutritionist_name']=person_name(getattr(item,'nutritionist_id',None))
            related_id=getattr(item,'patient_id',None) or getattr(item,'created_by',None) or getattr(item,'user_id',None)
            owner=people.get(related_id)
            branch_id=getattr(item,'branch_id',None)
            row.update(company_name=names.get(item.tenant_id,'Sin empresa'),branch_name=branch_names.get(branch_id,'Sin sucursal asignada'))
            rows.append({key:row.get(key,'') for key in selected})
        notes=['Datos limitados a las empresas, sucursales y pacientes autorizados.']
        if req.start_date or req.end_date:notes.append(f'Periodo en Bolivia (UTC−4): {req.start_date or "inicio"} a {req.end_date or "actualidad"}.')
        if req.status:notes.append('Estado: '+req.status+'.')
        for condition in req.filters:notes.append(f'Filtro: {condition.field} {condition.operator} {condition.value}.')
        if total>len(rows):notes.append(f'Mostrando {len(rows)} de {total} registros. Usa una agrupación para obtener totales completos.')
        return ReportQueryResponse(entity=entity_key,title=conf['label'],generated_at=datetime.now(timezone.utc),columns=[ReportColumnMeta(**all_cols[k]) for k in selected],rows=rows,total_rows=total,notes=notes)

    @classmethod
    def generate_excel(cls, report: ReportQueryResponse) -> io.BytesIO:
        """Exporta el reporte a un libro de trabajo Excel (.xlsx) con estilo profesional."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        import re
        clean_sheet_title = re.sub(r'[\\/*?:\[\]]', ' ', report.title).strip()[:31]
        ws.title = clean_sheet_title or "Reporte"

        # Colores institucionales NutriSalud (#206443)
        header_fill = PatternFill(start_color="206443", end_color="206443", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Segoe UI", size=14, bold=True, color="1B4D3E")
        meta_font = Font(name="Segoe UI", size=9, italic=True, color="666666")
        data_font = Font(name="Segoe UI", size=10, color="222222")
        zebra_fill = PatternFill(start_color="F7FAF8", end_color="F7FAF8", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='E0E0E0'),
            right=Side(style='thin', color='E0E0E0'),
            top=Side(style='thin', color='E0E0E0'),
            bottom=Side(style='thin', color='E0E0E0')
        )

        # 1. Título y metadatos
        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = f"NutriSalud - Reporte Dinámico: {report.title}"
        title_cell.font = title_font

        ws["A2"].value = f"Generado: {report.generated_at.strftime('%Y-%m-%d %H:%M:%S UTC')} | Total registros: {report.total_rows}"
        ws["A2"].font = meta_font

        if report.summary: ws['A3'].value=report.summary
        for index,note in enumerate(report.notes,start=4):ws.cell(row=index,column=1,value=note)
        # Context is included in exports alongside the exact result.
        start_row = 5+len(report.notes)
        for col_idx, col in enumerate(report.columns, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=col.label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[start_row].height = 24

        # 3. Filas de datos
        for row_idx, row_data in enumerate(report.rows, start=start_row + 1):
            is_zebra = (row_idx % 2 == 0)
            for col_idx, col in enumerate(report.columns, start=1):
                val = row_data.get(col.key, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                if isinstance(val,str):cell.data_type="s"
                cell.font = data_font
                cell.border = thin_border
                if is_zebra:
                    cell.fill = zebra_fill
                if col.type == "number":
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[row_idx].height = 20

        # Autoajustar ancho de columnas
        for col_idx, col in enumerate(report.columns, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(col.label)
            for r in report.rows:
                val_str = str(r.get(col.key, ""))
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 45))

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @classmethod
    def generate_pdf(cls, report: ReportQueryResponse) -> io.BytesIO:
        """Exporta el reporte a un documento PDF con tabla formateada e identidad visual."""
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        stream = io.BytesIO()
        # Usar orientación horizontal (landscape) para tablas dinámicas
        doc = SimpleDocTemplate(
            stream,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name="ReportTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            textColor=colors.HexColor("#206443"),
            spaceAfter=4,
        )
        meta_style = ParagraphStyle(
            name="ReportMeta",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            textColor=colors.HexColor("#555555"),
            spaceAfter=12,
        )
        table_header_style = ParagraphStyle(
            name="TableHeader",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9,
            textColor=colors.white,
            alignment=1,  # Center
        )
        table_cell_style = ParagraphStyle(
            name="TableCell",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            textColor=colors.HexColor("#222222"),
        )

        from xml.sax.saxutils import escape
        elements = []

        # Encabezado
        elements.append(Paragraph(escape(f"NutriSalud - Reporte Oficial: {report.title}"), title_style))
        elements.append(
            Paragraph(
                f"Emitido: {report.generated_at.strftime('%Y-%m-%d %H:%M UTC')} | "
                f"Total de registros: {report.total_rows} | Sistema Multi-Tenant NutriSalud",
                meta_style,
            )
        )
        if report.summary:elements.append(Paragraph(escape(report.summary),meta_style))
        for note in report.notes:elements.append(Paragraph(escape(note),meta_style))
        elements.append(Spacer(1, 8))

        # Construir matriz de datos para la tabla
        header_row = [Paragraph(c.label, table_header_style) for c in report.columns]
        table_data = [header_row]

        for row in report.rows:
            row_cells = []
            for col in report.columns:
                val = str(row.get(col.key, ""))
                # Truncar textos demasiado largos en PDF para evitar saltos masivos
                if len(val) > 70:
                    val = val[:67] + "..."
                row_cells.append(Paragraph(escape(val), table_cell_style))
            table_data.append(row_cells)

        # Crear tabla y aplicar estilos
        t = Table(table_data, repeatRows=1)
        t.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#206443")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAF8")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E0E0E0")),
                ("TOPPADDING", (0, 1), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
            ])
        )

        elements.append(t)
        doc.build(elements)

        stream.seek(0)
        return stream

    @classmethod
    def interpret_voice_command(
        cls,
        transcript: str,
        current_user: User,
        db: Session,
        tenant_id: Optional[str] = None,
        tenant_ids: Optional[List[str]] = None,
        branch_ids: Optional[List[str]] = None,
    ) -> VoiceReportCommandResponse:
        from app.services.report_intent_service import ReportIntentService
        return ReportIntentService.execute(cls, transcript, current_user, db, tenant_id=tenant_id, tenant_ids=tenant_ids, branch_ids=branch_ids)

    @classmethod
    def _heuristic_voice_parse(cls, transcript: str) -> Dict[str, Any]:
        t = transcript.lower()
        if any(w in t for w in ["cita", "agenda", "turno", "reserva", "consulta"]):
            entity = "appointments"
        elif any(w in t for w in ["receta", "comida", "alimento", "caloria", "proteina", "plato", "ingrediente"]):
            entity = "recipes"
        elif any(w in t for w in ["pago", "cobro", "dinero", "caja", "paypal", "ingreso", "transaccion"]):
            entity = "payments"
        elif any(w in t for w in ["clinica", "diagnostico", "ficha", "meta", "evolucion"]):
            entity = "clinical_records"
        elif any(w in t for w in ["bitacora", "log", "auditoria", "actividad", "movimiento"]):
            entity = "activity_logs"
        else:
            entity = "patients"

        status = None
        if entity == "appointments":
            if "confirmad" in t:
                status = "CONFIRMED"
            elif "pendient" in t:
                status = "PENDING"
            elif "cancelad" in t:
                status = "CANCELLED"
        elif entity == "patients":
            if "activ" in t:
                status = "Activo"
            elif "inactiv" in t:
                status = "Inactivo"
        elif entity == "payments":
            if "completad" in t or "pagad" in t:
                status = "COMPLETED"
            elif "pendient" in t:
                status = "PENDING"
        elif entity == "recipes":
            if "desayuno" in t:
                status = "Desayuno"
            elif "almuerzo" in t:
                status = "Almuerzo"
            elif "cena" in t:
                status = "Cena"
            elif "snack" in t:
                status = "Snack"

        label = cls.ENTITIES_CONFIG[entity]["label"]
        explanation = f"Comando por voz interpretado: Reporte de {label}"
        if status:
            explanation += f" con filtro de estado '{status}'"
        explanation += "."

        return {
            "entity": entity,
            "columns": None,
            "start_date": None,
            "end_date": None,
            "status": status,
            "search": None,
            "explanation": explanation,
        }

    @classmethod
    def generate_voice_summary(cls, req: VoiceReportSummaryRequest) -> VoiceReportSummaryResponse:
        """
        Genera un resumen analítico fluido con Gemini API redactado para ser leído en voz alta por TTS.
        """
        if GeminiService.is_configured():
            system_instruction = (
                "Eres el locutor y analista clínico de NutriSalud. "
                "Describe solo los números y hechos presentes. Usa la conclusión calculada y sus límites. No inventes recomendaciones clínicas ni sumes monedas diferentes. Una muestra de filas no representa el total; no extrapoles. Si no hay datos, dilo con claridad."
            )
            sample_data_str = json.dumps(req.sample_rows[:6], ensure_ascii=False)
            prompt = f"""
Analiza los datos del siguiente reporte de NutriSalud para redactar un resumen que será LEÍDO EN VOZ ALTA (locución de audio):
Reporte: {req.title}
Entidad: {req.entity}
Total de filas o grupos: {req.total_rows}
Conclusión calculada: {req.summary}
Criterios y límites: {req.notes}
Columnas: {req.columns}
Muestra de registros:
{sample_data_str}

Responde ÚNICAMENTE con un JSON con la siguiente estructura:
{{
  "summary_text": "Texto fluido y natural de 2 a 4 oraciones en español para ser locutado por voz alta (sintetiza el volumen de datos, métricas destacadas y hallazgos principales con tono profesional y positivo).",
  "bullet_points": [
    "Dato o hallazgo principal...",
    "Métrica relevante...",
    "Recomendación o estado..."
  ]
}}
"""
            try:
                result = GeminiService.generate_json(prompt, system_instruction=system_instruction)
                if result and isinstance(result, dict) and "summary_text" in result:
                    summary_text = result.get("summary_text", "")
                    bullet_points = result.get("bullet_points", [])
                    if isinstance(bullet_points, list) and bullet_points:
                        return VoiceReportSummaryResponse(
                            summary_text=summary_text,
                            bullet_points=[str(bp) for bp in bullet_points],
                        )
            except Exception:
                pass

        # Fallback si no está configurada la API de Gemini o si falla
        fallback_text = req.summary or f"Reporte de {req.title}. Se encontraron {req.total_rows} filas o grupos en el alcance consultado."
        fallback_points = ([req.summary] if req.summary else [f'Filas o grupos: {req.total_rows}']) + req.notes[:3]
        return VoiceReportSummaryResponse(summary_text=fallback_text,bullet_points=fallback_points)
